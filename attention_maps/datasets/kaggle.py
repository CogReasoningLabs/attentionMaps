"""Memory-bounded inspection and sampling for Kaggle-hosted datasets."""

from __future__ import annotations

import datetime as dt
import random
from pathlib import Path
from typing import Any, Iterator, Sequence


class KaggleDatasetError(ValueError):
    """Raised when a Kaggle dataset cannot be downloaded or read."""


def download_kaggle_file(dataset_id: str, dataset_file: str) -> Path:
    """Download one file through KaggleHub and return its cached local path."""

    try:
        import kagglehub
    except ImportError as error:
        raise KaggleDatasetError(
            "Kaggle datasets require `kagglehub`; install requirements.txt"
        ) from error

    try:
        downloaded = Path(
            kagglehub.dataset_download(dataset_id, path=dataset_file)
        ).expanduser()
    except Exception as error:
        raise KaggleDatasetError(
            f"Could not download Kaggle dataset {dataset_id!r}: {error}"
        ) from error

    if downloaded.is_file():
        return downloaded.resolve()

    candidate = downloaded / dataset_file
    if candidate.is_file():
        return candidate.resolve()

    matches = (
        list(downloaded.rglob(Path(dataset_file).name))
        if downloaded.is_dir()
        else []
    )
    if len(matches) == 1:
        return matches[0].resolve()
    raise KaggleDatasetError(
        f"Kaggle downloaded {dataset_id!r}, but {dataset_file!r} was not found"
    )


def _open_workbook(path: Path) -> Any:
    if path.suffix.lower() != ".xlsx":
        raise KaggleDatasetError(f"Unsupported Kaggle file type: {path.suffix}")
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise KaggleDatasetError(
            "Kaggle Excel datasets require `openpyxl`; install requirements.txt"
        ) from error
    try:
        return load_workbook(path, read_only=True, data_only=True)
    except Exception as error:
        raise KaggleDatasetError(f"Could not open Kaggle workbook: {error}") from error


def _column_names(values: Sequence[Any]) -> list[str]:
    names: list[str] = []
    occurrences: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        base = str(value).strip() if value is not None else ""
        base = base or f"column_{index}"
        occurrences[base] = occurrences.get(base, 0) + 1
        suffix = occurrences[base]
        names.append(base if suffix == 1 else f"{base}_{suffix}")
    return names


def _workbook_rows(workbook: Any) -> tuple[str, list[str], Iterator[tuple[Any, ...]]]:
    worksheet = workbook[workbook.sheetnames[0]]
    iterator = worksheet.iter_rows(values_only=True)
    for header in iterator:
        if any(value is not None for value in header):
            return worksheet.title, _column_names(header), iterator
    return worksheet.title, [], iter(())


def _value_type(value: Any) -> str:
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, float):
        return "number"
    if isinstance(value, dt.datetime):
        return "datetime"
    if isinstance(value, dt.date):
        return "date"
    if isinstance(value, dt.time):
        return "time"
    if isinstance(value, str):
        return "string"
    return type(value).__name__


def _serializable_value(value: Any) -> Any:
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    return value


def inspect_kaggle_workbook(dataset_id: str, dataset_file: str) -> dict[str, Any]:
    """Inspect the first worksheet without materializing all records."""

    path = download_kaggle_file(dataset_id, dataset_file)
    workbook = _open_workbook(path)
    try:
        sheet_name, columns, rows = _workbook_rows(workbook)
        observed_types = {column: set() for column in columns}
        nullable = {column: False for column in columns}
        row_count = 0
        for values in rows:
            values = tuple(values[: len(columns)])
            if not any(value is not None for value in values):
                continue
            row_count += 1
            for index, column in enumerate(columns):
                value = values[index] if index < len(values) else None
                if value is None:
                    nullable[column] = True
                else:
                    observed_types[column].add(_value_type(value))
    finally:
        workbook.close()

    return {
        "format": "kaggle",
        "dataset_id": dataset_id,
        "dataset_file": dataset_file,
        "path": str(path),
        "sheet_name": sheet_name,
        "rows": row_count,
        "files": 1,
        "bytes": path.stat().st_size,
        "row_groups": [],
        "columns": columns,
        "schema": [
            {
                "column": column,
                "type": " | ".join(sorted(observed_types[column])) or "null",
                "nullable": str(nullable[column]),
            }
            for column in columns
        ],
        "schema_variants": 1,
    }


def inspect_kaggle_text(dataset_id: str, dataset_file: str) -> dict[str, Any]:
    """Download and count a UTF-8 line corpus without retaining its records."""

    path = download_kaggle_file(dataset_id, dataset_file)
    if path.suffix.lower() != ".txt":
        raise KaggleDatasetError(f"Unsupported Kaggle text file type: {path.suffix}")
    try:
        with path.open("rb") as source:
            row_count = sum(1 for line in source if line.strip())
    except OSError as error:
        raise KaggleDatasetError(f"Could not scan Kaggle text file: {error}") from error

    return {
        "format": "kaggle_text",
        "dataset_id": dataset_id,
        "dataset_file": dataset_file,
        "path": str(path),
        "rows": row_count,
        "files": 1,
        "bytes": path.stat().st_size,
        "row_groups": [],
        "columns": ["text"],
        "schema": [{"column": "text", "type": "string", "nullable": "False"}],
        "schema_variants": 1,
    }


def sample_kaggle_workbook_rows(
    inventory: dict[str, Any],
    sample_size: int,
    seed: int,
    columns: Sequence[str],
) -> list[dict[str, Any]]:
    """Uniformly sample worksheet rows with an O(sample_size) reservoir."""

    selected_columns = list(columns)
    invalid = set(selected_columns) - set(inventory["columns"])
    if invalid:
        raise KaggleDatasetError(f"Unknown Kaggle columns: {sorted(invalid)}")
    if sample_size <= 0:
        return []

    path = Path(inventory["path"])
    workbook = _open_workbook(path)
    rng = random.Random(seed)
    reservoir: list[tuple[int, dict[str, Any]]] = []
    try:
        _, source_columns, rows = _workbook_rows(workbook)
        positions = {
            column: source_columns.index(column) for column in selected_columns
        }
        seen = 0
        for values in rows:
            values = tuple(values[: len(source_columns)])
            if not any(value is not None for value in values):
                continue
            record = {
                column: _serializable_value(
                    values[position] if position < len(values) else None
                )
                for column, position in positions.items()
            }
            record["__viewer_row_index"] = seen
            record["__viewer_file"] = (
                f"kaggle://datasets/{inventory['dataset_id']}/"
                f"{inventory['dataset_file']}"
            )
            seen += 1
            if len(reservoir) < sample_size:
                reservoir.append((seen - 1, record))
            else:
                replacement = rng.randrange(seen)
                if replacement < sample_size:
                    reservoir[replacement] = (seen - 1, record)
    finally:
        workbook.close()

    reservoir.sort(key=lambda item: item[0])
    return [record for _, record in reservoir]


def sample_kaggle_text_rows(
    inventory: dict[str, Any],
    sample_size: int,
    seed: int,
    columns: Sequence[str],
) -> list[dict[str, Any]]:
    """Read consecutive lines from a random byte offset with bounded memory."""

    selected_columns = list(columns)
    invalid = set(selected_columns) - {"text"}
    if invalid:
        raise KaggleDatasetError(f"Unknown Kaggle text columns: {sorted(invalid)}")
    requested_rows = min(sample_size, int(inventory["rows"]))
    if requested_rows <= 0:
        return []

    path = Path(inventory["path"])
    file_size = path.stat().st_size
    if file_size <= 0:
        return []
    start = random.Random(seed).randrange(file_size)
    records: list[dict[str, Any]] = []
    wrapped = False
    try:
        with path.open("rb") as source:
            source.seek(start)
            if start:
                source.readline()
            while len(records) < requested_rows:
                byte_offset = source.tell()
                line = source.readline()
                if not line:
                    if wrapped:
                        break
                    source.seek(0)
                    wrapped = True
                    continue
                text = line.decode("utf-8", errors="replace").strip()
                if not text:
                    continue
                record = {"text": text} if "text" in selected_columns else {}
                record["__viewer_row_index"] = f"byte:{byte_offset}"
                record["__viewer_file"] = (
                    f"kaggle://datasets/{inventory['dataset_id']}/"
                    f"{inventory['dataset_file']}"
                )
                records.append(record)
    except OSError as error:
        raise KaggleDatasetError(
            f"Could not sample Kaggle text file: {error}"
        ) from error
    return records
